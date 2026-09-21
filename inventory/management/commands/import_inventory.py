import hashlib
import io
import re
import urllib.request
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from inventory.models import ImportBatch, Location, Package, Reagent, StockMovement, Supplier

DEFAULT_SOURCE = "https://uapt33090-my.sharepoint.com/:x:/g/personal/lrocha_ua_pt/IQAWNeX4mo45S6fXA5kON1MQAeEI0trkb8UDvVU3lHDq-_Q?e=ttVMRJ&download=1"
HEADER_ALIASES = {
    "name": {"name", "nome", "garrafa", "gás", "gas"},
    "formula": {"formula", "fómula química", "fórmula química"},
    "supplier": {"supplier", "fornecedor"},
    "cas": {"cas", "cas number"},
    "quantity": {"quantity", "quantidade", "mass/volume"},
    "remaining": {"estimated quantity left", "quantity remaining"},
    "storage": {"storage", "localização", "sala/laboratório"},
    "shelf": {"shelf", "wood cabinet", "armário", "prateleira"},
    "sds": {"link to sds", "link to msds", "sds", "other sds"},
    "reference": {"reference", "referencia", "referência", "code"},
    "expiry": {"validade", "expired"},
    "physical_state": {"physical state"},
    "notes": {"notes", "notas"},
    "responsible": {"responsável", "responsavel"},
    "status": {"estado"},
}


def clean(value):
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def canonical_header(value):
    value = clean(value).lower().strip(":")
    if value.startswith("storage:"):
        return "storage"
    for canonical, aliases in HEADER_ALIASES.items():
        if value in aliases:
            return canonical
    return ""


def parse_quantity(raw):
    text = clean(raw).replace(",", ".")
    match = re.search(r"(?<![\d.])(\d+(?:\.\d+)?)\s*(µl|ml|cl|dl|l|mg|g|kg)\b", text, re.I)
    if not match:
        return None, ""
    try:
        return Decimal(match.group(1)), match.group(2).lower().replace("µ", "u")
    except InvalidOperation:
        return None, ""


def normalize_cas(raw):
    match = re.search(r"\b(\d{2,7}-\d{2}-\d)\b", clean(raw))
    if not match:
        return ""
    value = match.group(1)
    digits = value.replace("-", "")
    total = sum(int(digit) * index for index, digit in enumerate(reversed(digits[:-1]), start=1))
    return value if total % 10 == int(digits[-1]) else ""


def infer_state(value, name, sheet):
    text = f"{clean(value)} {name} {sheet}".lower()
    if "gas" in text or "gás" in text or "garrafa" in text:
        return Reagent.PhysicalState.GAS
    if "liquid" in text or "líquid" in text:
        return Reagent.PhysicalState.LIQUID
    if "solid" in text or "sólid" in text:
        return Reagent.PhysicalState.SOLID
    return Reagent.PhysicalState.UNKNOWN


class Command(BaseCommand):
    help = "Import the TEMA Excel inventory into normalized reagent and package records."

    def add_arguments(self, parser):
        parser.add_argument("--source", default=DEFAULT_SOURCE, help="XLSX path or public URL")
        parser.add_argument(
            "--force", action="store_true", help="Reprocess a previously imported workbook"
        )

    def read_source(self, source):
        if source.startswith(("http://", "https://")):
            try:
                request = urllib.request.Request(
                    source, headers={"User-Agent": "TEMA-Reagents/1.0"}
                )
                with urllib.request.urlopen(request, timeout=60) as response:
                    return response.read(), Path(response.url).name or "inventory.xlsx"
            except Exception as exc:
                raise CommandError(f"Could not download workbook: {exc}") from exc
        path = Path(source)
        if not path.exists():
            raise CommandError(f"Workbook not found: {path}")
        return path.read_bytes(), path.name

    @transaction.atomic
    def handle(self, *args, **options):
        content, source_name = self.read_source(options["source"])
        checksum = hashlib.sha256(content).hexdigest()
        existing = ImportBatch.objects.filter(checksum=checksum).first()
        if existing and not options["force"]:
            self.stdout.write(self.style.WARNING(f"Workbook already imported: {existing}"))
            return
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        skip_sheets = set()
        if "Inventário Gases (agosto 2026)" in workbook.sheetnames:
            skip_sheets.add("Inventário Gases (abril 2026)")

        building, _ = Location.objects.get_or_create(
            name="TEMA", code="TEMA", location_type=Location.Type.BUILDING, parent=None
        )
        imported = skipped = 0
        warnings = []

        for sheet in workbook.worksheets:
            if sheet.title in skip_sheets:
                warnings.append(f"{sheet.title}: skipped as an older snapshot")
                continue
            headers = {}
            sheet_lab = self.lab_code(sheet.title)
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                if not any(clean(value) for value in values):
                    continue
                mapped = {
                    canonical_header(value): index
                    for index, value in enumerate(values)
                    if canonical_header(value)
                }
                if "name" in mapped and len(mapped) >= 2:
                    headers = mapped
                    continue
                if not headers:
                    continue
                record = {
                    key: clean(values[index]) if index < len(values) else ""
                    for key, index in headers.items()
                }
                name = record.get("name", "")
                if not self.is_data_row(name, record):
                    skipped += 1
                    continue
                source_key = f"{checksum}:{sheet.title}:{row_number}"
                if Package.objects.filter(source_key=source_key).exists():
                    continue
                supplier = None
                if record.get("supplier"):
                    supplier, _ = Supplier.objects.get_or_create(name=record["supplier"][:200])
                cas = normalize_cas(record.get("cas", ""))
                reagent = self.get_reagent(name, cas, supplier, record, sheet.title, row_number)
                location = self.get_location(
                    building, sheet_lab, record.get("storage", ""), record.get("shelf", "")
                )
                quantity, unit = parse_quantity(record.get("quantity", ""))
                status = (
                    Package.Status.NOT_FOUND
                    if "não encontrei" in sheet.title.lower()
                    else Package.Status.AVAILABLE
                )
                notes = record.get("notes", "")
                if record.get("cas") and not cas:
                    warnings.append(
                        f"{sheet.title}:{row_number}: invalid/unrecognized CAS '{record['cas']}'"
                    )
                    notes = (notes + f" | Original CAS: {record['cas']}").strip(" |")
                package = Package.objects.create(
                    reagent=reagent,
                    initial_quantity=quantity,
                    current_quantity=quantity,
                    unit=unit,
                    raw_quantity=record.get("quantity", ""),
                    location=location,
                    status=status,
                    lot_number="",
                    notes=notes,
                    source_key=source_key,
                )
                StockMovement.objects.create(
                    package=package,
                    movement_type=StockMovement.Type.IMPORT,
                    quantity_after=quantity,
                    to_location=location,
                    reason=f"Imported from {sheet.title}, row {row_number}",
                )
                imported += 1

        if existing:
            existing.delete()
        batch = ImportBatch.objects.create(
            source_name=source_name[:300],
            checksum=checksum,
            imported_rows=imported,
            skipped_rows=skipped,
            warnings=warnings[:1000],
        )
        self.stdout.write(
            self.style.SUCCESS(
                f"Imported {imported} packages; skipped {skipped} rows; {len(warnings)} warnings. Batch {batch.pk}"
            )
        )

    def is_data_row(self, name, record):
        if not name or len(name) < 2:
            return False
        if canonical_header(name) or name.lower().startswith(
            ("inventário", "inventory", "atualização", "levantamento")
        ):
            return False
        return any(
            record.get(key)
            for key in ("cas", "formula", "supplier", "quantity", "reference", "storage")
        )

    def lab_code(self, text):
        match = re.search(r"\b\d+\.\d+\.\d+\b", text)
        return match.group(0) if match else "Unassigned"

    def get_reagent(self, name, cas, supplier, record, sheet, row):
        lookup = {"cas_number": cas} if cas else {"name__iexact": name, "supplier": supplier}
        reagent = Reagent.objects.filter(**lookup).first()
        if reagent:
            return reagent
        return Reagent.objects.create(
            name=name[:300],
            formula=record.get("formula", "")[:200],
            cas_number=cas,
            supplier=supplier,
            catalog_number=record.get("reference", "")[:120],
            physical_state=infer_state(record.get("physical_state", ""), name, sheet),
            default_unit=parse_quantity(record.get("quantity", ""))[1],
            sds_url=record.get("sds", "")
            if record.get("sds", "").startswith(("http://", "https://"))
            else "",
            description=""
            if record.get("sds", "").startswith(("http://", "https://"))
            else record.get("sds", ""),
            source_sheet=sheet,
            source_row=row,
        )

    def get_location(self, building, default_lab, storage, shelf):
        explicit_lab = self.lab_code(storage)
        lab_code = explicit_lab if explicit_lab != "Unassigned" else default_lab
        lab, _ = Location.objects.get_or_create(
            parent=building,
            code=lab_code,
            defaults={"name": f"Laboratory {lab_code}", "location_type": Location.Type.LABORATORY},
        )
        cabinet_name = (
            re.sub(r"\b\d+\.\d+\.\d+\b", "", clean(storage)).strip(" ()-,") or "Unspecified storage"
        )
        cabinet, _ = Location.objects.get_or_create(
            parent=lab,
            code=cabinet_name[:80],
            defaults={"name": cabinet_name[:200], "location_type": Location.Type.CABINET},
        )
        shelf_name = clean(shelf) or "Unspecified shelf"
        shelf_obj, _ = Location.objects.get_or_create(
            parent=cabinet,
            code=shelf_name[:80],
            defaults={"name": shelf_name[:200], "location_type": Location.Type.SHELF},
        )
        return shelf_obj
